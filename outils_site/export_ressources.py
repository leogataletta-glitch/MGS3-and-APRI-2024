"""Export for the non-tabbed sections ressources, apropos and contact.

Run from the site root:  python3 outils/export_ressources.py
Writes
  data/ressources/<lang>/0X_*.xlsx   the six published workbooks, built by the
                                     same functions as telechargements_page.py
  data/ressources/textes.json        every visible string (fr, en) of the three
                                     sections, read from the Streamlit modules
  data/ressources/apri-partage-<lang>.png   the "Share APRI" visual

Privacy: the workbooks only hold aggregates (counts, shares, scores by
group). The individual household dataset (d5) is never generated. In the
section x sub-population cross-classification (03) and the composite scores
built on it (04), any cell whose base is fewer than 5 households is
suppressed ("n < 5") or left out of the weighted mean.
"""
import io
import json
import os
import shutil
import sys

SITE = os.path.abspath(os.getcwd())
OUT = os.path.join(SITE, 'data', 'ressources')
sys.path.insert(0, '/tmp/work')
os.chdir('/tmp/work')

import streamlit as st  # noqa: E402
import i18n  # noqa: E402
import telechargements_page as tp  # noqa: E402
import a_propos_page as ap  # noqa: E402
import qualite_web as qw  # noqa: E402
import publication_web as pw  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

SEUIL = 5
_lire_orig = tp._lire_json
MODE = {'v': None}   # None | 'texte' | 'none'


def _lire_protege(nom):
    """ventilation.json with cells on fewer than 5 households masked."""
    d = _lire_orig(nom)
    if nom != 'ventilation.json' or MODE['v'] is None:
        return d
    d = json.loads(json.dumps(d))
    for blocs in d['sections'].values():
        for bloc in blocs.values():
            for p, n in (bloc.get('n') or {}).items():
                if n is not None and n < SEUIL:
                    masque = 'n < 5' if MODE['v'] == 'texte' else None
                    bloc['valeurs'][p] = masque
                    bloc['scores'][p] = masque
                    if MODE['v'] == 'texte':
                        bloc['n'][p] = '< 5'
    return d


tp._lire_json = _lire_protege


def lang(l):
    i18n.set_lang(l)
    st.session_state['lang'] = l


def ecrire_xlsx(l):
    lang(l)
    d = os.path.join(OUT, l)
    os.makedirs(d, exist_ok=True)
    jeux = [
        ('01_resultats_descriptifs.xlsx', tp._fichier_descriptif, None),
        ('02_indicateurs_resilience.xlsx', tp._fichier_indicateurs, None),
        ('03_ventilation_section_souspop.xlsx', tp._fichier_ventilation, 'texte'),
        ('04_scores_composites.xlsx', tp._fichier_composite, 'none'),
        ('06_dictionnaire_questionnaire.xlsx', tp._fichier_dictionnaire, None),
        ('07_organisations_communautaires.xlsx', tp._fichier_ocb, None),
    ]
    tailles = {}
    for nom, f, mode in jeux:
        MODE['v'] = mode
        fn = getattr(f, '__wrapped__', f)   # bypass st.cache_data
        data = fn(l)
        MODE['v'] = None
        with open(os.path.join(d, nom), 'wb') as fh:
            fh.write(data)
        tailles[nom] = len(data)
        wb = load_workbook(io.BytesIO(data), read_only=True)
        print(l, nom, len(data), 'bytes', wb.sheetnames)
    return tailles


STATIQUE = {
    'fr': ('Cette version de l’application est un site statique hébergé sur GitHub Pages ; elle ne dépose aucun cookie et n’intègre aucun outil de suivi d’usage. Cela ne couvre pas les journaux et services propres à l’hébergeur, aux polices de caractères ou aux fournisseurs de cartes. Signalez un problème d’accès ou de lecture en précisant la page, la langue et votre appareil.',
           'Audience : aucune statistique de fréquentation n’est collectée par le site lui-même. Aucun traceur publicitaire n’a été ajouté.'),
    'en': ('This version of the application is a static site hosted on GitHub Pages; it sets no cookie and includes no usage tracking. This does not cover the hosting provider’s, font providers’ or map providers’ own logs and services. Report access or readability issues with the page, language and device you used.',
           'Audience: the site itself collects no visit statistics. No advertising tracker has been added.'),
}


def textes(l):
    lang(l)
    T = i18n.T
    v = getattr(tp._volumes, '__wrapped__', tp._volumes)()
    nb = tp._nb
    jeux = []
    for cle, fichier, cv, n in [
            ('d1_titre', '01_resultats_descriptifs.xlsx', 'd_v_questions', v['q']),
            ('d2_titre', '02_indicateurs_resilience.xlsx', 'd_v_indicateurs', v['ind']),
            ('d3_titre', '03_ventilation_section_souspop.xlsx', 'd_v_indicateurs', v['ind']),
            ('d4_titre', '04_scores_composites.xlsx', 'd_v_dimensions', v['dim']),
            ('d6_titre', '06_dictionnaire_questionnaire.xlsx', 'd_v_items', v['q']),
            ('d7_titre', '07_organisations_communautaires.xlsx', 'd_v_organisations', v['ocb'])]:
        # same split as telechargements_page._bloc
        libelle = T(cle)
        if ' · ' in libelle:
            libelle = libelle.split(' · ', 1)[1]
        titre, _, sous = libelle.partition(', ')
        sous = sous[:1].upper() + sous[1:] if sous else ''
        volume = T(cv, n=nb(n))
        ch = ''.join(c for c in sous if c.isdigit())
        if ch and ch == ''.join(c for c in volume if c.isdigit()) and len(sous) <= 26:
            sous = ''
        jeux.append({'fichier': fichier, 'titre': titre, 'sous': sous,
                     'volume': volume, 'format': 'XLSX',
                     'description': T(cle.replace('_titre', '_desc'))})
    acces = qw.ACCESS[l]
    info = list(qw.TEXT[l])
    # The Streamlit wording describes Streamlit Community Cloud hosting and its
    # analytics; on the static site that would be false. Only the hosting
    # sentences change, the rest of the block is kept word for word.
    info[8] = STATIQUE[l][0]
    ap_t = {k: T(k) for k in ap.TEXTES}
    return {
        'ressources': {'jeux': jeux, 'bouton': T('d_bouton'),
                       'acces': {'titre': acces[0], 'texte': acces[1], 'bouton': acces[2]}},
        'apropos': {
            'lead': T('apx_lead'),
            'blocs': [{'k': T('apx_k%d' % i), 'x': T('apx_x%d' % i)} for i in (1, 2, 3)]
                     + [{'k': T('apx_k4')}],
            'paysages': [dict(zip(('nom', 'surface', 'detail'),
                                  (p.strip() for p in (T(c).split('|') + ['', '', ''])[:3])))
                         for c in ('ap_y_ga', 'ap_y_sud')],
            'jalons': [dict(zip(('v', 'l'), (p.strip() for p in (x.split('|') + [''])[:2])))
                       for x in T('apx_j').split('@@') if x.strip()],
            'realisations': {'titre': T('ap_o_t'),
                             'liste': [x.strip() for x in T('ap_o_l').split('@@') if x.strip()]},
            'definitions': {'titre': T('ap_d_t'),
                            'mots': [dict(zip(('mot', 'texte', 'source'),
                                              (p.strip() for p in (T(c).split('|') + ['', '', ''])[:3])))
                                     for c in ('ap_d_pay', 'ap_d_res')]},
            'informations': {
                'titre': info[0],
                'org': 'Programme des Nations Unies pour l’environnement (PNUE)' if l == 'fr'
                       else 'United Nations Environment Programme (UNEP)',
                'blocs': [{'t': info[j], 'x': info[j + 1]} for j in (1, 3, 5, 7)],
                'audience': STATIQUE[l][1],
            },
            'cgu': {
                'titre': pw.tr(pw.PAGES['cgu']),
                'maj': pw.tr(('Mise à jour : 21 septembre 2026', 'Updated: 21 September 2026', '', '')),
                'blocs': [{'t': h, 'x': x} for h, x in pw.TERMS[l]],
                'reference': pw.tr((
                    'Référence institutionnelle : [conditions d’utilisation du PNUE](https://www.unep.org/terms-use). Ces indications pratiques ne remplacent pas les conditions institutionnelles ni les droits propres aux sources.',
                    'Institutional reference: [UNEP terms of use](https://www.unep.org/terms-use). This practical guidance does not replace institutional terms or source-specific rights.', '', '')),
            },
            'plan': {'titre': pw.tr(('Plan du site', 'Site map', '', '')),
                     'pages': {k: pw.tr(v) for k, v in pw.PAGES.items()}},
            'partage': {'titre': pw.tr(('Partager APRI', 'Share APRI', '', '')),
                        'legende': pw.tr(('Visuel pour accompagner votre publication', 'Image to accompany your post', '', '')),
                        'image': 'apri-partage-%s.png' % l},
        },
        'contact': {k: ap_t[k] for k in ('ap_c_titre', 'ap_c_x', 'ap_c_mail', 'ap_c_qui',
                                         'ap_c_qui_x', 'ap_c_donnees', 'ap_c_donnees_x')},
        'courriel': ap.COURRIEL,
    }


def main():
    os.makedirs(OUT, exist_ok=True)
    tout = {}
    for l in ('fr', 'en'):
        ecrire_xlsx(l)
        tout[l] = textes(l)
        shutil.copy(os.path.join('/tmp/work/static', 'apri-partage-%s.png' % l), OUT)
    with open(os.path.join(OUT, 'textes.json'), 'w', encoding='utf-8') as f:
        json.dump(tout, f, ensure_ascii=False, indent=1)
    print('textes.json', os.path.getsize(os.path.join(OUT, 'textes.json')))


if __name__ == '__main__':
    main()
